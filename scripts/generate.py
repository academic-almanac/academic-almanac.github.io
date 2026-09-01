#!/usr/bin/env python3
"""
Academic Almanac - build data.json from content.xlsx.

    python scripts/generate.py                 # validate + write ./data.json
    python scripts/generate.py --dry-run       # validate only, write nothing
    python scripts/generate.py --keep-expired  # include past listings too
    python scripts/generate.py --excel PATH --out PATH

The workbook is the master: one sheet per category, headers matched by name
(column order may change, header text may not). Validation rules come from the
workbook's own "Legend & rules" and "Field Guide" sheets.

Errors block the write; warnings do not. Exit code 1 if anything blocked.
"""

import argparse
import json
import re
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required:  pip install openpyxl")


# --------------------------------------------------------------------------
# Schema. Excel header -> key in data.json, exactly as js/app.js reads them.
# --------------------------------------------------------------------------
FIELDS = {
    "id":                    "id",
    "category":              "cat",
    "acronym":               "acronym",
    "title":                 "title",
    "organiser":             "org",
    "city":                  "city",
    "country_code":          "country",
    "mode":                  "mode",
    "submission_deadline":   "subDeadline",
    "registration_deadline": "regDeadline",
    "feedback_date":         "feedbackDate",
    "venue_start":           "venueStart",
    "venue_end":             "venueEnd",
    "official_link":         "link",
    "submission_link":       "subLink",
    "description":           "body",
    "fee":                   "fee",
    "fee_max":               "feeMax",
    "currency":              "currency",
    "ects":                  "ects",
    "sector":                "sector",
    "author":                "author",
    "posted_date":           "ts",
    "verified":              "verified",
    "verified_date":         "verifiedDate",
}
RATING_COLS = ["rating_ABS", "rating_ABDC", "rating_VHB", "rating_FT50"]

DATE_COLS = ["submission_deadline", "registration_deadline", "feedback_date",
             "venue_start", "venue_end", "posted_date", "verified_date"]
NUM_COLS = ["fee", "fee_max", "ects"]

CATEGORIES = [
    "Conferences", "Symposiums", "PhD Courses & Summer Schools", "Workshops",
    "Calls for Papers", "Calls for Special Issues", "Grants & Funding",
    "Jobs & Positions", "Seminars & Webinars",
]
NON_CATEGORY_SHEETS = {"Legend & rules", "Field Guide"}

MODES = ["On-site", "Online", "Hybrid"]
# Extra modes allowed on specific sheets. A job can be genuinely remote, which
# is not the same claim as an event being held "Online".
MODES_EXTRA = {"Jobs & Positions": ["Remote"]}
SECTORS = ["Academic", "Private sector R&D", "Public sector & non-profit"]

# Required by the Field Guide, but a missing value is not worth blocking a
# build: js/app.js renders "posted by Unknown" when author is empty.
SOFT_REQUIRED = {"author"}
GRADES = {
    "rating_ABS":  ["4*", "4", "3", "2", "1"],
    "rating_ABDC": ["A*", "A", "B", "C"],
    "rating_VHB":  ["A+", "A", "B", "C", "D"],
    "rating_FT50": ["Listed"],
}

# The Field Guide sheet, transcribed. One character per category, in the order
# of CATEGORIES above:  X = required, x = optional, space = not used.
#                          Conf Symp PhD Wksp CfP CfSI Grant Jobs Semi
_GUIDE = {
    "acronym":               "xxxxxxxxx",
    "title":                 "XXXXXXXXX",
    "organiser":             "XXXXXXXXX",
    "city":                  "xxxxxxxxx",
    "country_code":          "xxxxxxxxx",
    "mode":                  "XXXXXXXXX",
    "submission_deadline":   "XXXXXXXX ",
    "registration_deadline": "xxxx    X",
    "feedback_date":         "xxxxxxxx ",
    "venue_start":           "XXXX    X",
    "venue_end":             "xxxx    x",
    "official_link":         "XXXXXXXXX",
    "submission_link":       "xxxxxxxxx",
    "description":           "XXXXXXXXX",
    "fee":                   "xxxxxxxXx",
    "fee_max":               "       x ",
    "currency":              "xxxxxxxXx",
    "ects":                  "  xx    x",
    "sector":                "       X ",
    "rating_ABS":            "xx xxx   ",
    "rating_ABDC":           "xx xxx   ",
    "rating_VHB":            "xx xxx   ",
    "rating_FT50":           "xx xxx   ",
    "author":                "XXXXXXXXX",
    "posted_date":           "XXXXXXXXX",
    "verified":              "XXXXXXXXX",
    "verified_date":         "xxxxxxxxx",
}


def guide(col, cat):
    """Return 'X' (required), 'x' (optional) or ' ' (not used for this cat)."""
    row = _GUIDE.get(col)
    if row is None:
        return "x"
    return row[CATEGORIES.index(cat)]


# --------------------------------------------------------------------------
# Cell coercion
# --------------------------------------------------------------------------
def as_text(v):
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def as_date(v, where, problems):
    """Excel date cell (or a parseable string) -> 'YYYY-MM-DD', or ''."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    for pat, order in ((r"^(\d{4})-(\d{1,2})-(\d{1,2})$", "ymd"),
                       (r"^(\d{1,2})[./](\d{1,2})[./](\d{4})$", "dmy")):
        m = re.match(pat, s)
        if m:
            a, b, c = m.groups()
            y, mo, d = (a, b, c) if order == "ymd" else (c, b, a)
            try:
                iso = date(int(y), int(mo), int(d)).isoformat()
            except ValueError:
                break
            problems.append(("warn", where, "stored as text " + repr(s) +
                             "; better to use a real Excel date cell"))
            return iso
    problems.append(("error", where, "cannot read " + repr(s) + " as a date "
                     "(use a real Excel date cell)"))
    return ""


def as_number(v, where, problems):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, bool):
        problems.append(("error", where, "expected a number, found TRUE/FALSE"))
        return None
    if isinstance(v, (int, float)):
        n = float(v)
    else:
        cleaned = re.sub(r"[^\d.\-]", "", str(v).strip())
        try:
            n = float(cleaned)
        except ValueError:
            problems.append(("error", where,
                             "expected a number, found " + repr(v)))
            return None
        problems.append(("warn", where, repr(v) + " read as " + format(n, "g") +
                         "; better to store it as a number"))
    return int(n) if n.is_integer() else n


def as_bool(v, where, problems):
    s = as_text(v).lower()
    if s in ("yes", "y", "true", "1"):
        return True
    if s in ("no", "n", "false", "0", ""):
        return False
    problems.append(("error", where,
                     "verified must be yes or no, found " + repr(v)))
    return False


# --------------------------------------------------------------------------
# Row -> post
# --------------------------------------------------------------------------
def build_post(raw, cat, sheet, excel_row, problems):
    """raw: {header: cell value}. Returns a post dict, appending problems."""
    def where(col):
        return sheet + "!row " + str(excel_row) + " [" + col + "]"

    post = {}
    for col, key in FIELDS.items():
        v = raw.get(col)
        if col in DATE_COLS:
            post[key] = as_date(v, where(col), problems)
        elif col in NUM_COLS:
            post[key] = as_number(v, where(col), problems)
        elif col == "verified":
            post[key] = as_bool(v, where(col), problems)
        else:
            post[key] = as_text(v)

    # id must be a whole number
    post["id"] = as_number(raw.get("id"), where("id"), problems)
    if post["id"] is None:
        problems.append(("error", where("id"), "required"))
    elif isinstance(post["id"], float):
        problems.append(("error", where("id"), "id must be a whole number"))

    # ratings: only the non-blank ones, keyed without the rating_ prefix
    ratings = {}
    for col in RATING_COLS:
        g = as_text(raw.get(col))
        if not g:
            continue
        system = col.replace("rating_", "")
        if guide(col, cat) == " ":
            problems.append(("warn", where(col), system + " rating is not used "
                             "for " + cat + "; kept, but check the entry"))
        elif g not in GRADES[col]:
            problems.append(("warn", where(col), repr(g) + " is not a " +
                             system + " grade (" + ", ".join(GRADES[col]) + ")"))
        ratings[system] = g
    post["ratings"] = ratings

    # ---- per-category required / not-used checks ----
    for col in list(FIELDS) + RATING_COLS:
        if col in ("id", "category"):
            continue
        need = guide(col, cat)
        filled = bool(as_text(raw.get(col)))
        if need == "X" and not filled:
            kind = "warn" if col in SOFT_REQUIRED else "error"
            problems.append((kind, where(col), "required for " + cat))
        elif need == " " and filled and col not in RATING_COLS:
            problems.append(("warn", where(col), "not used for " + cat +
                             "; kept, but check the entry"))

    # ---- value checks ----
    if post["cat"] and post["cat"] != cat:
        problems.append(("warn", where("category"), "says " + repr(post["cat"]) +
                         " but the sheet is " + repr(cat) +
                         "; using the sheet name"))
    post["cat"] = cat

    allowed_modes = MODES + MODES_EXTRA.get(cat, [])
    if post["mode"] and post["mode"] not in allowed_modes:
        problems.append(("error", where("mode"), "must be one of " +
                         " | ".join(allowed_modes) + ", found " +
                         repr(post["mode"])))
    if post["country"] and not re.fullmatch(r"[A-Za-z]{2}", post["country"]):
        problems.append(("warn", where("country_code"), repr(post["country"]) +
                         " is not a 2-letter ISO code"))
    post["country"] = post["country"].upper()
    if post["currency"] and not re.fullmatch(r"[A-Za-z]{3}", post["currency"]):
        problems.append(("warn", where("currency"), repr(post["currency"]) +
                         " is not a 3-letter ISO code"))
    post["currency"] = post["currency"].upper()
    if post["sector"] and post["sector"] not in SECTORS:
        problems.append(("warn", where("sector"), repr(post["sector"]) +
                         " is not one of " + " | ".join(SECTORS)))
    for key, col in (("link", "official_link"), ("subLink", "submission_link")):
        if post[key] and not post[key].lower().startswith(("http://", "https://")):
            problems.append(("error", where(col),
                             "must start with http:// or https://, found " +
                             repr(post[key])))
    if post["fee"] is not None and post["feeMax"] is not None \
            and post["feeMax"] < post["fee"]:
        problems.append(("warn", where("fee_max"), "fee_max is below fee"))
    if post["fee"] is not None and post["fee"] != 0 and not post["currency"]:
        problems.append(("warn", where("currency"),
                         "a fee is set but no currency"))
    if post["verified"] and not post["verifiedDate"]:
        problems.append(("warn", where("verified_date"),
                         "verified = yes but no verified_date"))
    if not post["verified"] and post["verifiedDate"]:
        problems.append(("warn", where("verified_date"),
                         "verified_date is set but verified = no"))

    # dates should run in a sensible order
    order = [("submission_deadline", "subDeadline"),
             ("feedback_date", "feedbackDate"),
             ("venue_start", "venueStart"),
             ("venue_end", "venueEnd")]
    seen = [(c, post[k]) for c, k in order if post[k]]
    for (col1, d1), (col2, d2) in zip(seen, seen[1:]):
        if d2 < d1:
            problems.append(("warn", where(col2),
                             d2 + " is before " + col1 + " (" + d1 + ")"))

    return post


def expiry_of(post):
    """Last date this listing is still worth showing. '' if it has no dates."""
    dates = [d for d in (post["subDeadline"], post["regDeadline"],
                         post["venueStart"], post["venueEnd"]) if d]
    return max(dates) if dates else ""


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------
def main():
    root = Path(__file__).resolve().parent.parent      # the site / repo root
    ap = argparse.ArgumentParser(
        description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--excel", type=Path,
                    default=root.parent / "Excel4GitHub" / "content.xlsx",
                    help="master workbook (default: ../Excel4GitHub/content.xlsx)")
    ap.add_argument("--out", type=Path, default=root / "data.json",
                    help="output file (default: ./data.json)")
    ap.add_argument("--dry-run", action="store_true",
                    help="validate and report, write nothing")
    ap.add_argument("--keep-expired", action="store_true",
                    help="include listings whose dates have all passed")
    ap.add_argument("--force", action="store_true",
                    help="write even if there are errors (not recommended)")
    args = ap.parse_args()

    if not args.excel.exists():
        sys.exit("Workbook not found: " + str(args.excel))

    wb = openpyxl.load_workbook(args.excel, data_only=True, read_only=True)
    problems = []
    posts, expired = [], []
    today = date.today().isoformat()
    known = set(FIELDS) | set(RATING_COLS)

    for sheet in wb.sheetnames:
        if sheet in NON_CATEGORY_SHEETS:
            continue
        if sheet not in CATEGORIES:
            problems.append(("warn", sheet, "not a known category sheet - skipped"))
            continue

        rows = wb[sheet].iter_rows(values_only=True)
        try:
            header = [as_text(h) for h in next(rows)]
        except StopIteration:
            problems.append(("warn", sheet, "sheet is empty - skipped"))
            continue

        missing = known - set(header)
        if missing:
            problems.append(("error", sheet, "missing column(s): " +
                             ", ".join(sorted(missing))))
            continue
        for h in header:
            if h and h not in known:
                problems.append(("warn", sheet + " [" + h + "]",
                                 "unknown column - ignored"))

        n = 0
        for excel_row, row in enumerate(rows, start=2):
            if not any(as_text(c) for c in row):
                continue                                   # blank spacer row
            post = build_post(dict(zip(header, row)), sheet, sheet,
                              excel_row, problems)
            exp = expiry_of(post)
            if exp and exp < today and not args.keep_expired:
                expired.append(post)
            else:
                posts.append(post)
            n += 1
        if n == 0:
            problems.append(("info", sheet, "no data rows - this category will "
                             "not appear on the site"))

    # ids must be unique across the whole workbook, expired rows included
    counts = Counter(p["id"] for p in posts + expired if p["id"] is not None)
    for i in sorted(i for i, c in counts.items() if c > 1):
        problems.append(("error", "workbook", "id " + str(i) +
                         " is used more than once"))

    # nearest submission deadline first; undated last
    posts.sort(key=lambda p: (p["subDeadline"] == "", p["subDeadline"], p["title"]))

    # ---- report ----
    errors = [p for p in problems if p[0] == "error"]
    warns = [p for p in problems if p[0] == "warn"]
    infos = [p for p in problems if p[0] == "info"]
    for kind, where, msg in errors + warns + infos:
        print(kind.upper().ljust(5) + " " + where + ": " + msg, file=sys.stderr)

    by_cat = Counter(p["cat"] for p in posts)
    line = "\n" + str(len(posts)) + " listing(s) in " + str(len(by_cat))
    line += " category" if len(by_cat) == 1 else " categories"
    if expired:
        line += ", " + str(len(expired)) + " expired and dropped"
    print(line)
    for cat in CATEGORIES:
        if by_cat[cat]:
            print("  " + str(by_cat[cat]).rjust(4) + "  " + cat)
    print("  " + str(len(errors)) + " error(s), " + str(len(warns)) +
          " warning(s)")

    if errors and not args.force:
        print("\nNothing written - fix the errors above, or pass --force.",
              file=sys.stderr)
        return 1
    if args.dry_run:
        print("\nDry run - nothing written.")
        return 1 if errors else 0

    payload = {
        "generatedAt": datetime.now().replace(microsecond=0).isoformat(),
        "posts": posts,
    }
    args.out.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n",
                        encoding="utf-8")
    print("\nWrote " + str(args.out) + " (" +
          format(args.out.stat().st_size, ",") + " bytes)")
    print("Commit data.json and push to publish.")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
